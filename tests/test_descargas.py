"""Centro de Descargas: los tres reportes se generan de verdad y traen lo correcto.

Se ejecuta el mismo código que usa el botón "Descargar reporte"; solo se
reemplazan los diálogos (dónde guardar, avisos) para que no se abra nada.
"""
import csv
import os
import shutil
import tempfile
from unittest import mock

from base import PruebaVIM
from stub_tk import modulo_reportes

try:
    import openpyxl
except ImportError:  # el Excel no se puede generar sin openpyxl
    openpyxl = None


class _BaseDescargas(PruebaVIM):
    def setUp(self):
        super().setUp()
        self.carpeta = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.carpeta, True)

        from database.ventas import anular_venta, confirmar_venta
        self.caja = self.abrir_caja()
        self.configurar_iva(19, incluido=True)
        self.pid = self.crear_producto(nombre="Bebida", precio_venta=1190, precio_compra=500, stock=50)
        self.ventas = []
        for cantidad, metodo in ((1, "tarjeta"), (2, "efectivo"), (3, "efectivo")):
            ok, msg, vid = confirmar_venta(
                [self.linea_ticket(self.pid, nombre="Bebida", cantidad=cantidad, precio_unitario=1190)],
                100000, metodo, usuario_id=1)
            self.assertTrue(ok, msg)
            self.ventas.append(vid)
        ok, msg = anular_venta(self.ventas[0], 1, "prueba")
        self.assertTrue(ok, msg)

        self.hoy = self.ahora_local()[:10]
        self.hoy_dmy = "-".join(reversed(self.hoy.split("-")))

        vm = modulo_reportes()
        self.pantalla = vm.ModuloVentasReportes.__new__(vm.ModuloVentasReportes)
        self.pantalla.root = None
        self.listos, self.avisos = [], []
        self.pantalla._avisar_reporte_listo = lambda ruta, *a, **k: self.listos.append(ruta)

    def _descargar(self, nombre, accion):
        """Ejecuta una descarga guardando en `nombre`, sin abrir diálogos."""
        import tkinter.filedialog
        import tkinter.messagebox
        ruta = os.path.join(self.carpeta, nombre)
        registrar = (lambda tipo: lambda *a, **k: self.avisos.append((tipo, a)))
        with mock.patch.object(tkinter.filedialog, "asksaveasfilename", lambda *a, **k: ruta, create=True), \
             mock.patch.object(tkinter.messagebox, "showerror", registrar("error"), create=True), \
             mock.patch.object(tkinter.messagebox, "showwarning", registrar("aviso"), create=True), \
             mock.patch.object(tkinter.messagebox, "showinfo", registrar("info"), create=True):
            accion()
        self.assertEqual(self.avisos, [], f"La descarga mostró un mensaje: {self.avisos}")
        self.assertEqual(self.listos, [ruta])
        self.assertTrue(os.path.exists(ruta))
        return ruta


class PruebaCsvVentas(_BaseDescargas):
    def test_lista_todas_las_ventas_y_suma_solo_las_completadas(self):
        ruta = self._descargar("ventas.csv", lambda: self.pantalla._ejecutar_exportacion_especifica(
            self.hoy_dmy, self.hoy_dmy, self.hoy, self.hoy, 2))
        with open(ruta, encoding="utf-8-sig", newline="") as f:
            filas = list(csv.reader(f, delimiter=";"))

        ventas = [f for f in filas if f and f[0].startswith("VTA-")]
        self.assertEqual(len(ventas), 3)
        self.assertEqual(sorted(f[-1] for f in ventas), ["Anulada", "Completada", "Completada"])
        totales = next(f for f in filas if f and f[0].startswith("TOTALES"))
        self.assertEqual(totales[8], "5950", "La venta anulada no debe sumar")

    def test_periodo_sin_ventas_avisa_y_no_crea_archivo(self):
        import tkinter.messagebox
        avisos = []
        with mock.patch.object(tkinter.messagebox, "showwarning", lambda *a, **k: avisos.append(a), create=True):
            self.pantalla._ejecutar_exportacion_especifica("01-01-2001", "01-01-2001", "2001-01-01", "2001-01-01", 2)
        self.assertEqual(len(avisos), 1)
        self.assertEqual(os.listdir(self.carpeta), [])


@__import__("unittest").skipIf(openpyxl is None, "openpyxl no está instalado")
class PruebaExcelVentasYResultados(_BaseDescargas):
    def test_trae_ventas_y_estado_de_resultados(self):
        ruta = self._descargar("ventas.xlsx", lambda: self.pantalla._ejecutar_exportacion_especifica(
            self.hoy_dmy, self.hoy_dmy, self.hoy, self.hoy, 1))
        wb = openpyxl.load_workbook(ruta)
        self.assertEqual(wb.sheetnames, ["Ventas Detalladas", "Estado de Resultados Financiero"])

        pnl = {fila[0].strip(): fila[2] for fila in wb["Estado de Resultados Financiero"].iter_rows(values_only=True)
               if fila and isinstance(fila[0], str) and fila[2] is not None}
        self.assertEqual(pnl["(+) Ventas Netas"], 5000)       # 5.950 con IVA
        self.assertEqual(pnl["(i) Impuestos (IVA)"], 950)
        self.assertEqual(pnl["(-) Costo de Ventas (COGS)"], 2500)
        self.assertEqual(pnl["Punto de equilibrio (ventas necesarias)"],
                         round(pnl["Punto de equilibrio (ventas necesarias)"]), "Debe ir sin decimales")


@__import__("unittest").skipIf(openpyxl is None, "openpyxl no está instalado")
class PruebaReporteMaestro(_BaseDescargas):
    def _maestro(self):
        ruta = self._descargar("maestro.xlsx", lambda: self.pantalla._ejecutar_super_reporte(
            self.hoy, self.hoy, self.hoy_dmy, self.hoy_dmy))
        return openpyxl.load_workbook(ruta)

    def test_trae_las_cinco_hojas(self):
        self.assertEqual(self._maestro().sheetnames,
                         ["Ventas Detalladas", "Movimientos de Caja", "Inventario",
                          "Cuentas por Pagar", "Nómina Personal"])

    def test_ventas_detalladas_excluye_anuladas(self):
        filas = list(self._maestro()["Ventas Detalladas"].iter_rows(values_only=True))
        detalle = [f for f in filas if f[0] and str(f[0]).startswith("VTA-")]
        self.assertEqual(len(detalle), 2)
        total = next(f for f in filas if f[0] == "TOTAL")
        self.assertEqual(total[12], 5950)

    def test_inventario_valoriza_productos_a_granel(self):
        """Los productos por kilo tienen precio_venta 0: antes salían valorizados en $0."""
        with self.conectar() as con:
            con.execute("""
                INSERT INTO productos (nombre, tipo_item, precio_compra, precio_venta, precio_por_kilo,
                                       stock, stock_minimo, tipo_venta, activo, fecha_creacion)
                VALUES ('Tomate granel', 'producto', 900, 0, 1500, 10, 0, 'peso', 1, ?)
            """, (self.ahora_local(),))
            con.commit()
        filas = list(self._maestro()["Inventario"].iter_rows(values_only=True))
        tomate = next(f for f in filas if f[0] == "Tomate granel")
        self.assertEqual(tomate[7], 1500)       # precio por kilo
        self.assertEqual(tomate[9], 15000)      # 10 kg × $1.500


class PruebaAnulacionEnCaja(PruebaVIM):
    """Anular una venta con tarjeta no puede sacar billetes del cajón."""

    def _resumen(self, caja):
        from database.caja import obtener_resumen_caja
        return obtener_resumen_caja(caja)

    def _vender_y_anular(self, metodo):
        from database.ventas import anular_venta, confirmar_venta
        caja = self.abrir_caja(monto=10000)
        self.configurar_iva(0)
        pid = self.crear_producto(precio_venta=3000, stock=10)
        ok, msg, vid = confirmar_venta([self.linea_ticket(pid, precio_unitario=3000)], 3000, metodo, usuario_id=1)
        self.assertTrue(ok, msg)
        self.assertTrue(anular_venta(vid, 1, "prueba")[0])
        return self._resumen(caja)

    def test_anular_tarjeta_no_cambia_el_efectivo(self):
        r = self._vender_y_anular("tarjeta")
        self.assertEqual(r["efectivo_esperado"], 10000)
        self.assertEqual(r["ventas_tarjeta"], 0)
        self.assertEqual(r["egresos"], 0)

    def test_anular_transferencia_no_cambia_el_efectivo(self):
        r = self._vender_y_anular("transferencia")
        self.assertEqual(r["efectivo_esperado"], 10000)
        self.assertEqual(r["ventas_transferencia"], 0)

    def test_anular_efectivo_devuelve_la_plata_del_cajon(self):
        r = self._vender_y_anular("efectivo")
        self.assertEqual(r["efectivo_esperado"], 10000)   # 10.000 + 3.000 − 3.000
        self.assertEqual(r["egresos"], 3000)
